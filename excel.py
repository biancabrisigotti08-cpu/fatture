"""
Costruzione del file Excel: fogli Fatture, Righe e Duplicati.
"""
import io
import re
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLORE_INTESTAZIONE = "FF6B35"

# Testi dell'Excel in italiano e in inglese
TESTI = {
    "it": {
        "fogli": ("Fatture", "Righe", "Duplicati"),
        "fatture": ["Fornitore", "P.IVA fornitore", "Cliente", "P.IVA cliente", "Tipo documento",
                    "Numero", "Data", "Imponibile (€)", "IVA (€)", "Totale (€)", "Valuta",
                    "N. righe", "File", "Letto da"],
        "righe": ["Fornitore", "Numero", "Data", "Descrizione", "Quantità",
                  "Prezzo unitario (€)", "Aliquota IVA %", "Importo (€)"],
        "extra": {"targa": "Targa", "telaio": "Telaio", "categoria": "Categoria"},
        "gia_presente": "Già presente nel file",
        "file": "estrazione_fatture.xlsx",
    },
    "en": {
        "fogli": ("Invoices", "Line items", "Duplicates"),
        "fatture": ["Supplier", "Supplier VAT no.", "Customer", "Customer VAT/Tax ID", "Document type",
                    "Number", "Date", "Taxable amount", "VAT", "Total", "Currency",
                    "Lines", "File", "Read from"],
        "righe": ["Supplier", "Number", "Date", "Description", "Quantity",
                  "Unit price", "VAT rate %", "Amount"],
        "extra": {"targa": "Licence plate", "telaio": "VIN", "categoria": "Category"},
        "gia_presente": "Already in file",
        "file": "invoices_extracted.xlsx",
    },
}

TIPI_EN = {
    "Fattura": "Invoice", "Acconto su fattura": "Advance invoice", "Acconto su parcella": "Advance fee note",
    "Nota di credito": "Credit note", "Nota di debito": "Debit note", "Parcella": "Fee note",
    "Integrazione reverse charge": "Reverse charge integration", "Autofattura servizi estero": "Self-invoice (foreign services)",
    "Integrazione beni UE": "EU goods integration", "Integrazione beni art.17": "Goods integration (art. 17)",
    "Autofattura regolarizzazione": "Self-invoice (regularisation)", "Autofattura splafonamento": "Self-invoice (ceiling exceeded)",
    "Estrazione deposito IVA": "VAT warehouse release", "Estrazione deposito IVA con versamento": "VAT warehouse release with payment",
    "Fattura differita": "Deferred invoice", "Fattura differita triangolare": "Deferred triangular invoice",
    "Cessione beni ammortizzabili": "Sale of depreciable assets", "Autoconsumo / omaggio": "Self-consumption / gift",
    "Acquisti da San Marino": "Purchases from San Marino", "Omessa o irregolare fatturazione": "Missing or irregular invoicing",
    "Ricevuta": "Receipt",
}
LETTO_EN = {
    "XML": "XML", "XML firmato (.p7m)": "Signed XML (.p7m)", "PDF con XML allegato": "PDF with embedded XML",
    "PDF (AI)": "PDF (AI)", "PDF (lettura base)": "PDF (basic reading)",
}
CATEGORIE_EN = {"Km Eccedenti": "Excess mileage", "Danni": "Damage", "Perizia": "Appraisal",
                "Tagliando": "Service", "ETM": "Missing technical items"}


def _tr(diz, valore, lingua):
    return diz.get(valore, valore) if lingua == "en" else valore


def nome_file(lingua="it"):
    return TESTI.get(lingua, TESTI["it"])["file"]
FORMATO_EURO = '#,##0.00'
FORMATO_DATA = 'DD/MM/YYYY'


def categoria(descrizione):
    """Categorie usate per le fatture dei veicoli (colonna mostrata solo se serve)."""
    d = (descrizione or "").lower()
    for chiave, nome in (("forfait", "Forfait"), ("over plafond", "Over Plafond"),
                         ("km eccedenti", "Km Eccedenti"), ("esubero km", "Km Eccedenti"),
                         ("eccedenza chilometrica", "Km Eccedenti"), ("tagliando", "Tagliando"),
                         ("perizia", "Perizia"), ("elementi tecnici mancanti", "ETM"),
                         ("penale per danni", "Danni")):
        if chiave in d:
            return nome
    if re.search(r'\beam\b', d):
        return "EAM"
    return ""


def _data(val):
    try:
        a, m, g = (int(x) for x in val.split('-'))
        return date(a, m, g)
    except Exception:
        return val or None


def _chiave(f):
    """Due fatture sono la stessa se coincidono fornitore, numero, data e totale."""
    numero = re.sub(r'\s', '', (f.get("numero") or "")).upper()
    if not numero:
        return None
    fornitore = (f.get("piva_fornitore") or f.get("fornitore") or "").upper().replace(" ", "")
    totale = round(f["totale"], 2) if isinstance(f.get("totale"), float) else None
    return (fornitore, numero, f.get("data") or "", totale)


def _intestazione(ws, colonne):
    riempimento = PatternFill("solid", fgColor=COLORE_INTESTAZIONE)
    for c, (titolo, larghezza, _) in enumerate(colonne, start=1):
        cella = ws.cell(row=1, column=c, value=titolo)
        cella.fill = riempimento
        cella.font = Font(bold=True, color="FFFFFF")
        cella.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = larghezza
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _scrivi(ws, colonne, righe):
    for valori in righe:
        ws.append(valori)
    for c, (_, _, formato) in enumerate(colonne, start=1):
        if formato:
            for (cella,) in ws.iter_rows(min_row=2, min_col=c, max_col=c):
                cella.number_format = formato
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonne))}{ws.max_row}"


_FORMATI_FATTURE = [(32, None), (16, None), (30, None), (16, None), (16, None), (16, None),
                    (12, FORMATO_DATA), (15, FORMATO_EURO), (13, FORMATO_EURO), (15, FORMATO_EURO),
                    (8, None), (9, None), (30, None), (18, None)]


def colonne_fatture(lingua):
    return [(t, w, fmt) for t, (w, fmt) in zip(TESTI[lingua]["fatture"], _FORMATI_FATTURE)]


def _riga_fattura(f, lingua="it"):
    return [f["fornitore"], f["piva_fornitore"], f["cliente"], f["piva_cliente"],
            _tr(TIPI_EN, f["tipo"], lingua), f["numero"], _data(f["data"]), f["imponibile"], f["iva"],
            f["totale"], f["valuta"], len(f["righe"]), f["file"], _tr(LETTO_EN, f["letto_con"], lingua)]


def crea_excel(fatture, lingua="it"):
    """Restituisce (BytesIO, n_fatture_uniche, n_duplicati)."""
    lingua = lingua if lingua in TESTI else "it"
    T = TESTI[lingua]
    col_fatture = colonne_fatture(lingua)
    uniche, duplicati, visti = [], [], {}
    for f in fatture:
        k = _chiave(f)
        if k and k in visti:
            duplicati.append((f, visti[k]))
        else:
            if k:
                visti[k] = f["file"]
            uniche.append(f)

    wb = openpyxl.Workbook()

    # Foglio 1: una riga per fattura
    ws = wb.active
    ws.title = T["fogli"][0]
    _intestazione(ws, col_fatture)
    _scrivi(ws, col_fatture, [_riga_fattura(f, lingua) for f in uniche])

    # Foglio 2: una riga per ogni voce delle fatture
    righe = [(f, r) for f in uniche for r in f["righe"]]
    con_targa = any(r["targa"] for _, r in righe)
    con_telaio = any(r["telaio"] for _, r in righe)
    con_categoria = any(categoria(r["descrizione"]) for _, r in righe)
    formati = [(30, None), (14, None), (12, FORMATO_DATA), (50, None), (10, '#,##0.##'),
               (16, FORMATO_EURO), (12, '0.##'), (15, FORMATO_EURO)]
    colonne = [(t, w, fmt) for t, (w, fmt) in zip(T["righe"], formati)]
    if con_targa:
        colonne.append((T["extra"]["targa"], 13, None))
    if con_telaio:
        colonne.append((T["extra"]["telaio"], 21, None))
    if con_categoria:
        colonne.append((T["extra"]["categoria"], 16, None))
    ws2 = wb.create_sheet(T["fogli"][1])
    _intestazione(ws2, colonne)
    dati_righe = []
    for f, r in righe:
        valori = [f["fornitore"], f["numero"], _data(f["data"]), r["descrizione"],
                  r["quantita"], r["prezzo_unitario"], r["aliquota_iva"], r["importo"]]
        if con_targa:
            valori.append(r["targa"])
        if con_telaio:
            valori.append(r["telaio"])
        if con_categoria:
            valori.append(_tr(CATEGORIE_EN, categoria(r["descrizione"]), lingua))
        dati_righe.append(valori)
    _scrivi(ws2, colonne, dati_righe)

    # Foglio 3: fatture caricate più di una volta
    colonne_dup = col_fatture + [(T["gia_presente"], 30, None)]
    ws3 = wb.create_sheet(T["fogli"][2])
    _intestazione(ws3, colonne_dup)
    _scrivi(ws3, colonne_dup, [_riga_fattura(f, lingua) + [primo] for f, primo in duplicati])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out, len(uniche), len(duplicati)
